#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 Excel 导入测试 fixtures（可重复运行，产物直接提交到 git）。

产物（desktop/rust/tests/fixtures/）：
  excel_multi.xlsx  3 个可见工作表 + 1 个隐藏工作表，含数字/日期/时间/布尔/中文/NULL
  excel_multi.xls   同内容的 BIFF8（Excel 97-2003）工作簿，3 个工作表

两张工作簿的前三张表内容完全一致，方便 Rust 测试用同一组断言覆盖
ZIP(xlsx) 与 OLE2(xls) 两条读取路径：

  员工（第一张，默认选中）：id / 姓名 / 入职日期 / 活跃 / 评分 / 备注 / 登录时间 / 打卡时间
  订单：订单号 / 客户 / 金额
  配置：键 / 值

复现命令（使用 DSH 自带 Python；xlwt 不在默认 site-packages 里）：

  PY=/Users/lyp/.dsh/dsh-runtimes/dsh-next-primary-runtime/dependencies/python/bin/python3
  $PY -m pip install --quiet --target /tmp/xlwt_pkg xlwt
  PYTHONPATH=/tmp/xlwt_pkg $PY gen_fixtures.py

openpyxl 写 .xlsx（ZIP + OOXML），xlwt 写 .xls（OLE2 + BIFF8）。
日期/时间单元格靠「数值 + 数字格式」表达：calamine 的 dates 特性把有日期格式的
数值读成 ExcelDateTime；纯时间（<1 的数值）会落在 1899-12-31，用来测 time 推断。
"""

from __future__ import annotations

import datetime
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- 测试数据
# None = 空单元格（NULL）。日期/时间用 datetime 对象，布尔用 Python bool。
EMPLOYEES_HEADER = ["id", "姓名", "入职日期", "活跃", "评分", "备注", "登录时间", "打卡时间"]
EMPLOYEES = [
    [1, "张伟", datetime.date(2021, 3, 5), True, 88.5, None,
     datetime.datetime(2021, 3, 5, 9, 30, 0), datetime.time(9, 30, 0)],
    [2, "李娜", datetime.date(2022, 7, 18), False, 92, "优秀",
     datetime.datetime(2022, 7, 18, 14, 5, 30), datetime.time(9, 5, 30)],
    [3, "王强", datetime.date(2023, 1, 2), True, 76.25, None,
     datetime.datetime(2023, 1, 2, 8, 0, 0), datetime.time(8, 45, 0)],
    [4, "赵敏", datetime.date(2023, 11, 30), True, 66, "试用",
     datetime.datetime(2023, 11, 30, 18, 20, 0), datetime.time(10, 15, 0)],
]
ORDERS_HEADER = ["订单号", "客户", "金额"]
ORDERS = [
    [1001, "甲公司", 299.5],
    [1002, "乙公司", 1280],
    [1003, "丙公司", 58.75],
]
CONFIG_HEADER = ["键", "值"]
CONFIG = [
    ["version", "1.0"],
    ["env", "prod"],
]


def is_time_only(value: object) -> bool:
    return isinstance(value, datetime.time) and not isinstance(value, datetime.datetime)


# ---------------------------------------------------------------- xlsx
def write_xlsx(path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "员工"
    ws.append(EMPLOYEES_HEADER)
    for row in EMPLOYEES:
        ws.append(row)
    # openpyxl 会按 Python 类型自动写布尔/数字；日期/时间需要显式数字格式，
    # 否则 calamine 只会看到一串裸数值。
    for index, row in enumerate(EMPLOYEES, start=2):
        ws.cell(row=index, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=index, column=7).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=index, column=8).number_format = "hh:mm:ss"

    orders = wb.create_sheet("订单")
    orders.append(ORDERS_HEADER)
    for row in ORDERS:
        orders.append(row)

    config = wb.create_sheet("配置")
    config.append(CONFIG_HEADER)
    for row in CONFIG:
        config.append(row)

    hidden = wb.create_sheet("隐藏数据")
    hidden.append(["secret"])
    hidden.append(["不应出现在工作表列表"])
    hidden.sheet_state = "hidden"

    wb.save(path)


# ---------------------------------------------------------------- xls
def write_xls(path: str) -> None:
    try:
        import xlwt
    except ImportError:  # pragma: no cover - 复现命令写在模块 docstring 里
        sys.exit(
            "缺少 xlwt：PYTHONPATH=/tmp/xlwt_pkg 后重跑（见本文件顶部复现命令）"
        )

    wb = xlwt.Workbook(encoding="utf-8")
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    datetime_style = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
    time_style = xlwt.easyxf(num_format_str="HH:MM:SS")

    def add_sheet(name: str, header: list[str], rows: list[list]) -> None:
        sheet = wb.add_sheet(name)
        for column, title in enumerate(header):
            sheet.write(0, column, title)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                if value is None:
                    continue  # 空单元格 = NULL
                if is_time_only(value):
                    # 时间写成一天内的比例，配合 HH:MM:SS 格式；BIFF 没有单独的 time 类型。
                    seconds = (
                        value.hour * 3600 + value.minute * 60 + value.second
                    )
                    sheet.write(r, c, seconds / 86400.0, time_style)
                elif isinstance(value, datetime.datetime):
                    sheet.write(r, c, value, datetime_style)
                elif isinstance(value, datetime.date):
                    sheet.write(r, c, datetime.datetime.combine(value, datetime.time()), date_style)
                else:
                    sheet.write(r, c, value)

    add_sheet("员工", EMPLOYEES_HEADER, EMPLOYEES)
    add_sheet("订单", ORDERS_HEADER, ORDERS)
    add_sheet("配置", CONFIG_HEADER, CONFIG)
    wb.save(path)


def main() -> None:
    xlsx = os.path.join(HERE, "excel_multi.xlsx")
    xls = os.path.join(HERE, "excel_multi.xls")
    write_xlsx(xlsx)
    write_xls(xls)
    for path in (xlsx, xls):
        print(f"wrote {os.path.basename(path)} ({os.path.getsize(path)} bytes)")


if __name__ == "__main__":
    main()
