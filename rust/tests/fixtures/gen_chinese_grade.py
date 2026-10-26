#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成中文测试表格（第 50 轮追加）：学生名单 + 年级成绩单。

工作表 1「学生名单」：姓名 / 学号 / 班级
工作表 2「年级成绩单」：学号 / 姓名 / 班级 / 语文 / 数学 / 英语 / 科学 / 体育（每科整年级成绩）

用 DSH 自带 Python（openpyxl）生成：
  PY=python3
  $PY gen_chinese_grade.py
"""
from __future__ import annotations

import os
import random

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "chinese_grade.xlsx")

random.seed(20261024)

SURNAMES = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
GIVEN = ["伟", "芳", "娜", "敏", "静", "强", "磊", "洋", "艳", "勇", "军", "杰", "娟", "涛", "明",
         "超", "秀英", "霞", "平", "刚", "桂英", "文", "辉", "丽", "鹏", "晨", "雪", "宁", "健", "梅"]
CLASSES = ["高一(1)班", "高一(2)班", "高一(3)班"]
SUBJECTS = ["语文", "数学", "英语", "科学", "体育"]

STUDENTS = []
for i in range(30):
    name = random.choice(SURNAMES) + random.choice(GIVEN)
    student_no = "2026%02d%02d" % (i // 10 + 1, i % 10 + 1)
    klass = CLASSES[i // 10]
    STUDENTS.append((name, student_no, klass))

ROWS = []
for name, student_no, klass in STUDENTS:
    scores = []
    for subject in SUBJECTS:
        # 缺考用空单元格（NULL）表示，导入时应保持 NULL。
        scores.append(None if random.random() < 0.03 else random.randint(55, 100))
    ROWS.append([student_no, name, klass, *scores])


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet(title="学生名单")
    ws1.append(["姓名", "学号", "班级"])
    for name, student_no, klass in STUDENTS:
        ws1.append([name, student_no, klass])
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet(title="年级成绩单")
    ws2.append(["学号", "姓名", "班级", *SUBJECTS])
    for row in ROWS:
        ws2.append(row)
    ws2.freeze_panes = "A2"

    wb.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes", "rows:", len(STUDENTS), len(ROWS))


if __name__ == "__main__":
    main()
